"""
scripts/build_results_matrix.py

Builds the paper-facing Excel results matrix requested by the team:
one row per executed experiment (or experiment "block"), columns
(blocco di esperimento, path, risultato, RQ a cui risponde, validato).

Why this is a separate, offline script rather than something inside
eiger/experiments/
----------------------------------------------------------------------
This is a reporting/bookkeeping artifact for the paper, not part of the
measurement pipeline itself, and depends on openpyxl (an optional,
data-import-only dependency — see pyproject.toml). Keeping it here
mirrors scripts/enrich_snopes_claims.py and
scripts/clean_snopes_contamination.py: standalone, not covered by the
100% coverage gate (see pyproject.toml's [tool.coverage.run] omit list).

What this script does
-----------------------
1. Globs every ``results/**/results.json`` produced by
   ``eiger.experiments.ExperimentRunner`` (see ``eiger/experiments/README.md``
   for the schema: ``experiment_id``, ``config``, ``records``,
   ``aggregate_metrics``, ``environment``).
2. For each one, builds a one-line "risultato" summary from
   ``aggregate_metrics`` (rounded to 4 decimals) plus ``n_records``.
3. Maps ``experiment_id`` to a research question via
   ``_EXPERIMENT_TO_RQ`` below — a hand-maintained table, NOT inferred
   automatically, because only a human can correctly say which RQ a
   given experiment design actually answers. Add new experiment_ids to
   this table as they're defined; an unmapped experiment_id is reported
   as "RQ non assegnata — aggiornare _EXPERIMENT_TO_RQ" rather than
   guessed, so the matrix never silently misattributes a result.
4. "validato" is always "No" on first generation — this script has no
   way to know whether a human has reviewed a given experiment's
   results, and defaulting to "No" is the conservative direction for a
   column that will be read by paper co-authors. Once a row has been
   reviewed, a human edits that cell in the generated .xlsx directly;
   re-running this script from scratch will reset it, so treat the
   .xlsx as something to hand-edit after generation, not regenerate
   blindly once review has started (or extend this script to read back
   prior "validato" values by experiment_id — not implemented here).
5. Writes ``results/results_matrix.xlsx`` (path overridable via -o).

Usage
------
    pip install 'eiger[data-import]'   # one-time: installs openpyxl
    python scripts/build_results_matrix.py
    python scripts/build_results_matrix.py -o /tmp/matrix.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required for this script. Install it with: "
        "pip install 'eiger[data-import]' (or: pip install openpyxl)"
    ) from exc

_DEFAULT_RESULTS_DIR = Path("results")
_DEFAULT_OUTPUT = Path("results/results_matrix.xlsx")

# Hand-maintained experiment_id -> RQ mapping. Update this whenever a new
# experiments/*.yaml is added or an existing one's purpose changes — see
# docs/CLAIM_AND_RESEARCH_QUESTIONS.md for the RQ1-RQ5 definitions.
_EXPERIMENT_TO_RQ: dict[str, str] = {
    "baseline_v1": "Baseline / condizione di controllo (poison_rate=0.0) — riferimento per RQ1-RQ5, non risponde direttamente a nessuna RQ",
    "snopes_pilot": "Validazione dataset (SnopesDataset end-to-end) — non una RQ, pilot tecnico",
    "ablation_attacks": "RQ1 (vulnerabilita per tipo di manipolazione) / RQ2 (tasso di poisoning) — verificare contro la config esatta prima di pubblicare",
}

_HEADERS = [
    "Blocco di esperimento",
    "Path",
    "Risultato (aggregate_metrics)",
    "RQ a cui risponde",
    "Validato",
]


def _find_results_files(results_dir: Path) -> list[Path]:
    return sorted(results_dir.glob("**/results.json"))


def _format_result(data: dict[str, Any]) -> str:
    agg = data.get("aggregate_metrics", {})
    n_records = len(data.get("records", []))
    metrics_str = ", ".join(f"{k}={v:.4f}" for k, v in agg.items())
    return f"n={n_records} | {metrics_str}" if metrics_str else f"n={n_records} | (nessuna metrica)"


def _rq_for(experiment_id: str) -> str:
    return _EXPERIMENT_TO_RQ.get(
        experiment_id,
        "RQ non assegnata — aggiornare _EXPERIMENT_TO_RQ in scripts/build_results_matrix.py",
    )


def build_rows(results_dir: Path) -> list[list[str]]:
    """Build one row per results.json found under results_dir."""
    rows: list[list[str]] = []
    for path in _find_results_files(results_dir):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as exc:
            rows.append([str(path.parent.name), str(path), f"ERRORE lettura file: {exc}", "N/D", "No"])
            continue

        experiment_id = data.get("experiment_id", path.parent.name)
        rows.append(
            [
                experiment_id,
                str(path),
                _format_result(data),
                _rq_for(experiment_id),
                "No",
            ]
        )
    return rows


def write_matrix(rows: list[list[str]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results Matrix"

    ws.append(_HEADERS)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_idx in range(1, len(_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        ws.append(row)

    for row_idx in range(2, len(rows) + 2):
        for col_idx in range(1, len(_HEADERS) + 1):
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    widths = [22, 45, 55, 60, 12]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "--results-dir",
        type=Path,
        default=_DEFAULT_RESULTS_DIR,
        help=f"Directory to search for results.json files (default: {_DEFAULT_RESULTS_DIR})",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=_DEFAULT_OUTPUT,
        help=f"Output .xlsx path (default: {_DEFAULT_OUTPUT})",
    )
    args = parser.parse_args(argv)

    rows = build_rows(args.results_dir)
    if not rows:
        print(f"No results.json files found under '{args.results_dir}'. Nothing to write.", file=sys.stderr)
        return 1

    write_matrix(rows, args.output)
    print(f"Wrote {len(rows)} row(s) to {args.output}")
    for row in rows:
        print(f"  - {row[0]}: {row[2]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
