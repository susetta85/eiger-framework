"""
Unit tests for scripts/compute_calibration_correlation.py.

Not covered by the coverage gate (pyproject.toml omits scripts/*), but the
statistics helpers (Pearson/Spearman/MAE, rank computation with ties) and
the annotation/reference-file join logic are non-trivial enough to warrant
regression tests.
"""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pytest

_SCRIPT_PATH = (
    Path(__file__).resolve().parents[2] / "scripts" / "compute_calibration_correlation.py"
)


@pytest.fixture(scope="module")
def corr_module():
    spec = importlib.util.spec_from_file_location("compute_calibration_correlation", _SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class TestRescale:
    def test_rescales_1_to_0_and_5_to_1(self, corr_module) -> None:
        assert corr_module._rescale_1_to_5(1.0) == 0.0
        assert corr_module._rescale_1_to_5(5.0) == 1.0
        assert corr_module._rescale_1_to_5(3.0) == pytest.approx(0.5)


class TestPearsonR:
    def test_perfect_positive_correlation(self, corr_module) -> None:
        assert corr_module.pearson_r([1, 2, 3, 4], [10, 20, 30, 40]) == pytest.approx(1.0)

    def test_perfect_negative_correlation(self, corr_module) -> None:
        assert corr_module.pearson_r([1, 2, 3, 4], [40, 30, 20, 10]) == pytest.approx(-1.0)

    def test_fewer_than_two_points_returns_zero(self, corr_module) -> None:
        assert corr_module.pearson_r([1.0], [2.0]) == 0.0
        assert corr_module.pearson_r([], []) == 0.0

    def test_zero_variance_returns_zero_not_raising(self, corr_module) -> None:
        """A constant sequence has zero variance; must degrade gracefully, not divide by zero."""
        assert corr_module.pearson_r([1, 1, 1], [1, 2, 3]) == 0.0


class TestRanks:
    def test_no_ties_gives_sequential_ranks(self, corr_module) -> None:
        assert corr_module._ranks([30, 10, 20]) == [3.0, 1.0, 2.0]

    def test_ties_get_averaged_ranks(self, corr_module) -> None:
        assert corr_module._ranks([10, 20, 20, 30]) == [1.0, 2.5, 2.5, 4.0]

    def test_all_tied_gives_identical_average_rank(self, corr_module) -> None:
        assert corr_module._ranks([5, 5, 5]) == [2.0, 2.0, 2.0]


class TestSpearmanRho:
    def test_perfect_monotonic_relationship(self, corr_module) -> None:
        assert corr_module.spearman_rho([1, 2, 3], [10, 100, 1000]) == pytest.approx(1.0)


class TestMeanAbsoluteError:
    def test_computes_mean_of_absolute_differences(self, corr_module) -> None:
        assert corr_module.mean_absolute_error([1.0, 2.0], [1.5, 1.0]) == pytest.approx(0.75)

    def test_empty_input_returns_zero(self, corr_module) -> None:
        assert corr_module.mean_absolute_error([], []) == 0.0


class TestInterpret:
    def test_weak_moderate_strong_bands(self, corr_module) -> None:
        assert corr_module._interpret(0.1) == "weak"
        assert corr_module._interpret(0.5) == "moderate"
        assert corr_module._interpret(0.9) == "strong"
        assert corr_module._interpret(-0.9) == "strong"  # magnitude, not sign


class TestLoadAnnotations:
    def _write_workbook(self, path, rows) -> None:
        import openpyxl

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append(["claim_id", "query", "faithfulness_human", "correctness_human"])
        for row in rows:
            sheet.append(row)
        wb.save(path)

    def test_skips_rows_with_blank_human_scores(self, corr_module, tmp_path) -> None:
        path = tmp_path / "batch.xlsx"
        self._write_workbook(path, [
            ["C1", "q1", 4, 5],
            ["C2", "q2", None, None],  # not yet annotated
            ["C3", "q3", 3, None],  # partially annotated
        ])
        annotations = corr_module.load_annotations(path)
        assert [a["claim_id"] for a in annotations] == ["C1"]

    def test_missing_required_column_raises(self, corr_module, tmp_path) -> None:
        import openpyxl

        path = tmp_path / "batch.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["claim_id", "query"])
        wb.save(path)
        with pytest.raises(ValueError, match="faithfulness_human"):
            corr_module.load_annotations(path)


class TestMainCLI:
    def test_missing_workbook_returns_1(self, corr_module, tmp_path, capsys) -> None:
        exit_code = corr_module.main([str(tmp_path / "nope.xlsx")])
        assert exit_code == 1
        assert "not found" in capsys.readouterr().err

    def test_missing_reference_file_returns_1(self, corr_module, tmp_path, capsys) -> None:
        import openpyxl

        workbook_path = tmp_path / "batch.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["claim_id", "faithfulness_human", "correctness_human"])
        wb.save(workbook_path)

        exit_code = corr_module.main([str(workbook_path)])
        assert exit_code == 1
        assert "not found" in capsys.readouterr().err

    def test_end_to_end_prints_correlation(self, corr_module, tmp_path, capsys) -> None:
        import openpyxl

        workbook_path = tmp_path / "batch.xlsx"
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append(["claim_id", "faithfulness_human", "correctness_human"])
        sheet.append(["C1", 5, 1])
        sheet.append(["C2", 1, 5])
        wb.save(workbook_path)

        reference_path = workbook_path.with_suffix(workbook_path.suffix + ".automated_scores.json")
        reference_path.write_text(json.dumps({
            "source_results_json": "results/e/results.json",
            "faithfulness_scorer_config": "embedding",
            "scores": {
                "C1": {"ragas_faithfulness": 1.0, "ragas_answer_correctness": 0.0},
                "C2": {"ragas_faithfulness": 0.0, "ragas_answer_correctness": 1.0},
            },
        }))

        exit_code = corr_module.main([str(workbook_path)])
        out = capsys.readouterr().out

        assert exit_code == 0
        assert "Faithfulness" in out
        assert "Answer correctness" in out
        assert "Pearson r  = +1.000" in out  # C1 human=1.0/auto=1.0, C2 human=0.0/auto=0.0

    def test_no_annotated_rows_returns_0(self, corr_module, tmp_path, capsys) -> None:
        import openpyxl

        workbook_path = tmp_path / "batch.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["claim_id", "faithfulness_human", "correctness_human"])
        wb.active.append(["C1", None, None])
        wb.save(workbook_path)

        reference_path = workbook_path.with_suffix(workbook_path.suffix + ".automated_scores.json")
        reference_path.write_text(json.dumps({"source_results_json": "x", "faithfulness_scorer_config": "embedding", "scores": {}}))

        exit_code = corr_module.main([str(workbook_path)])
        assert exit_code == 0
        assert "No fully-annotated rows" in capsys.readouterr().out

    def test_claim_id_with_no_automated_score_is_skipped_not_crashing(self, corr_module, tmp_path, capsys) -> None:
        import openpyxl

        workbook_path = tmp_path / "batch.xlsx"
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append(["claim_id", "faithfulness_human", "correctness_human"])
        sheet.append(["C1", 5, 5])
        sheet.append(["UNKNOWN_ID", 3, 3])
        wb.save(workbook_path)

        reference_path = workbook_path.with_suffix(workbook_path.suffix + ".automated_scores.json")
        reference_path.write_text(json.dumps({
            "source_results_json": "x", "faithfulness_scorer_config": "embedding",
            "scores": {"C1": {"ragas_faithfulness": 1.0, "ragas_answer_correctness": 1.0}},
        }))

        exit_code = corr_module.main([str(workbook_path)])
        out = capsys.readouterr().out
        assert exit_code == 0
        assert "1 had no matching automated score" in out
