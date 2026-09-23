"""
Unit tests for CorpusClaimDataset (eiger.datasets.corpus_claim).

Covers the schema mapping from the collaborator's ``01_Corpus_claim``
sheet to Claim, and — most importantly — the ethical gate that makes
load() return zero claims by default (every row currently has
requires_human_review=True; see the class's own module docstring).

What these tests do NOT cover:
  - The real Corpus_claim_RAG_Mistral_output_v3.xlsx workbook itself (it
    is not committed to this repository — a private, human-curated
    corpus). A small synthetic workbook is built per test instead.
"""

from __future__ import annotations

import sys
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

from eiger.core.exceptions import IngestionError
from eiger.datasets import CorpusClaimDataset, get_dataset, list_datasets

_HEADER = (
    "item_id", "source_platform", "factcheck_url", "claim_original",
    "original_rating", "normalized_label", "topic", "subtopic", "risk_level",
    "sensitivity_class", "review_status", "requires_human_review",
    "preflight_status", "generation_status", "modified_claim",
    "modified_verdict_normalized", "manipulation_type_applied",
    "changed_element", "claim_change_description", "pipeline_safety_flags",
    "original_language_iso",
)


@pytest.fixture(autouse=True)
def _silence_logger() -> Iterator[None]:
    with patch("eiger.datasets.corpus_claim.log"):
        yield


@pytest.fixture()
def workbook_path(tmp_path: Path) -> Path:
    return tmp_path / "Corpus_claim_RAG_Mistral_output_v3.xlsx"


def _row(
    item_id: str = "SNP_000002",
    claim_original: str = "Whoopi Goldberg's mansion was destroyed in the 2025 wildfires.",
    normalized_label: str = "verified_false",
    risk_level: int = 4,
    sensitivity_class: str = "S1",
    requires_human_review: bool = True,
    preflight_status: str | None = None,
    **overrides: Any,
) -> dict[str, Any]:
    row: dict[str, Any] = {
        "item_id": item_id,
        "source_platform": "Snopes",
        "factcheck_url": "https://example.org/fact-check",
        "claim_original": claim_original,
        "original_rating": "false",
        "normalized_label": normalized_label,
        "topic": "Politics",
        "subtopic": "Elections",
        "risk_level": risk_level,
        "sensitivity_class": sensitivity_class,
        "review_status": "Approvata",
        "requires_human_review": requires_human_review,
        "preflight_status": preflight_status,
        "generation_status": "generated",
        "modified_claim": "A modified variant of the claim.",
        "modified_verdict_normalized": "misleading",
        "manipulation_type_applied": "M02",
        "changed_element": "Year",
        "claim_change_description": "The year was changed.",
        "pipeline_safety_flags": None,
        "original_language_iso": "en",
    }
    row.update(overrides)
    return row


def _write_workbook(path: Path, rows: list[dict[str, Any]]) -> None:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "01_Corpus_claim"
    sheet.append(_HEADER)
    for row in rows:
        sheet.append([row.get(col) for col in _HEADER])
    wb.save(path)


# ─── Identity ─────────────────────────────────────────────────────────────────

class TestCorpusClaimDatasetIdentity:
    def test_name_and_description(self) -> None:
        dataset = CorpusClaimDataset(path="/tmp/whatever.xlsx")
        assert dataset.name == "corpus_claim"
        assert "mistral" in dataset.description.lower()

    def test_default_path_points_at_data_corpus_claim(self) -> None:
        dataset = CorpusClaimDataset()
        assert dataset.path.parts[-3:] == (
            "data", "corpus_claim", "Corpus_claim_RAG_Mistral_output_v3.xlsx",
        )

    def test_custom_path_override(self, workbook_path: Path) -> None:
        dataset = CorpusClaimDataset(path=workbook_path)
        assert dataset.path == workbook_path

    def test_registered_under_corpus_claim_name(self) -> None:
        assert "corpus_claim" in list_datasets()
        assert isinstance(get_dataset("corpus_claim"), CorpusClaimDataset)


# ─── The ethical gate ─────────────────────────────────────────────────────────

class TestCorpusClaimDatasetEthicalGate:
    def test_load_returns_empty_by_default_when_all_rows_unreviewed(
        self, workbook_path: Path
    ) -> None:
        _write_workbook(workbook_path, [_row(requires_human_review=True)])
        claims = CorpusClaimDataset(path=workbook_path).load()
        assert claims == []

    def test_load_includes_reviewed_rows_by_default(self, workbook_path: Path) -> None:
        _write_workbook(workbook_path, [_row(requires_human_review=False)])
        claims = CorpusClaimDataset(path=workbook_path).load()
        assert len(claims) == 1

    def test_include_unreviewed_returns_all_non_blocked_rows(self, workbook_path: Path) -> None:
        rows = [
            _row(item_id="SNP_1", requires_human_review=True),
            _row(item_id="SNP_2", requires_human_review=False),
        ]
        _write_workbook(workbook_path, rows)
        claims = CorpusClaimDataset(path=workbook_path).load(include_unreviewed=True)
        assert {c.claim_id for c in claims} == {"SNP_1", "SNP_2"}

    def test_returned_claims_still_carry_requires_human_review_metadata(
        self, workbook_path: Path
    ) -> None:
        _write_workbook(workbook_path, [_row(requires_human_review=False)])
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert claim.metadata["requires_human_review"] is False

    def test_blocked_rows_excluded_even_with_include_unreviewed(
        self, workbook_path: Path
    ) -> None:
        rows = [
            _row(item_id="SNP_1", requires_human_review=True, preflight_status="blocked"),
            _row(item_id="SNP_2", requires_human_review=True),
        ]
        _write_workbook(workbook_path, rows)
        claims = CorpusClaimDataset(path=workbook_path).load(include_unreviewed=True)
        assert [c.claim_id for c in claims] == ["SNP_2"]

    def test_s3_rows_excluded_even_with_include_unreviewed(self, workbook_path: Path) -> None:
        """
        docs/ETHICS_AND_THREAT_MODEL.md §5: S3 ("excluded") must never
        enter the corpus at all, unlike S0-S2. No S3 rows exist in the
        real file inspected while writing this loader, but the schema
        allows the value, so this loader must not silently let one through.
        """
        rows = [
            _row(item_id="SNP_1", requires_human_review=True, sensitivity_class="S3"),
            _row(item_id="SNP_2", requires_human_review=True, sensitivity_class="S2"),
        ]
        _write_workbook(workbook_path, rows)
        claims = CorpusClaimDataset(path=workbook_path).load(include_unreviewed=True)
        assert [c.claim_id for c in claims] == ["SNP_2"]

    def test_blank_row_is_skipped_not_loaded_as_none_text(self, workbook_path: Path) -> None:
        """
        Regression test: a blank trailing row (common in real spreadsheet
        exports) used to become a Claim with claim_id="None" and
        original_fact="None" instead of being skipped, since both cells
        are simply absent (None) rather than triggering a missing-column
        KeyError.
        """
        rows = [
            {col: None for col in _HEADER},
            _row(item_id="SNP_1", requires_human_review=False),
        ]
        _write_workbook(workbook_path, rows)
        claims = CorpusClaimDataset(path=workbook_path).load()
        assert [c.claim_id for c in claims] == ["SNP_1"]


# ─── Field mapping ────────────────────────────────────────────────────────────

class TestCorpusClaimDatasetFieldMapping:
    def test_maps_claim_id_and_text(self, workbook_path: Path) -> None:
        _write_workbook(workbook_path, [_row(item_id="SNP_000042", requires_human_review=False)])
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert claim.claim_id == "SNP_000042"
        assert claim.original_fact.startswith("Whoopi Goldberg")

    def test_source_dataset_is_corpus_claim(self, workbook_path: Path) -> None:
        _write_workbook(workbook_path, [_row(requires_human_review=False)])
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert claim.source_dataset == "corpus_claim"

    def test_templated_context_query(self, workbook_path: Path) -> None:
        _write_workbook(
            workbook_path,
            [_row(claim_original="Cats can fly.", requires_human_review=False)],
        )
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert claim.context_query == "Is it true that Cats can fly.?"

    def test_ground_truth_label_maps_directly(self, workbook_path: Path) -> None:
        rows = [
            _row(item_id="SNP_1", normalized_label="verified_true", requires_human_review=False),
            _row(item_id="SNP_2", normalized_label="verified_false", requires_human_review=False),
        ]
        _write_workbook(workbook_path, rows)
        claims = CorpusClaimDataset(path=workbook_path).load()
        by_id = {c.claim_id: c.ground_truth_label for c in claims}
        assert by_id == {"SNP_1": "verified_true", "SNP_2": "verified_false"}

    def test_unrecognized_normalized_label_maps_to_none(self, workbook_path: Path) -> None:
        _write_workbook(
            workbook_path,
            [_row(normalized_label="something_else", requires_human_review=False)],
        )
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert claim.ground_truth_label is None

    def test_risk_level_and_sensitivity_class_map_directly(self, workbook_path: Path) -> None:
        _write_workbook(
            workbook_path,
            [_row(risk_level=5, sensitivity_class="S2", requires_human_review=False)],
        )
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert claim.risk_level == 5
        assert claim.sensitivity_class == "S2"

    def test_provenance_metadata_carried_when_present(self, workbook_path: Path) -> None:
        _write_workbook(workbook_path, [_row(requires_human_review=False)])
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert claim.metadata["manipulation_type_applied"] == "M02"
        assert claim.metadata["modified_claim"] == "A modified variant of the claim."
        assert claim.metadata["topic"] == "Politics"

    def test_none_provenance_fields_omitted_from_metadata(self, workbook_path: Path) -> None:
        _write_workbook(
            workbook_path,
            [_row(requires_human_review=False, pipeline_safety_flags=None)],
        )
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert "pipeline_safety_flags" not in claim.metadata

    def test_original_language_iso_carried_into_metadata(self, workbook_path: Path) -> None:
        """
        Bug fix regression test: original_language_iso was read from the
        workbook but never added to _PROVENANCE_METADATA_COLUMNS, so it was
        silently dropped at load time — no downstream code could tell a
        non-English claim (36/5,672 rows in the real corpus, 17 genuinely
        Spanish) apart from an English one.
        """
        _write_workbook(
            workbook_path,
            [_row(requires_human_review=False, original_language_iso="es")],
        )
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert claim.metadata["original_language_iso"] == "es"

    def test_none_original_language_iso_omitted_from_metadata(self, workbook_path: Path) -> None:
        _write_workbook(
            workbook_path,
            [_row(requires_human_review=False, original_language_iso=None)],
        )
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert "original_language_iso" not in claim.metadata

    def test_invalid_sensitivity_class_degrades_to_none_without_raising(
        self, workbook_path: Path
    ) -> None:
        """
        This is an external, human-maintained spreadsheet — a malformed
        cell (e.g. a typo'd "S4") must not raise a pydantic ValidationError
        that aborts loading every other, valid row in the same load() call.
        """
        rows = [
            _row(item_id="SNP_1", sensitivity_class="S4", requires_human_review=False),
            _row(item_id="SNP_2", sensitivity_class="S1", requires_human_review=False),
        ]
        _write_workbook(workbook_path, rows)
        claims = CorpusClaimDataset(path=workbook_path).load()
        by_id = {c.claim_id: c.sensitivity_class for c in claims}
        assert by_id == {"SNP_1": None, "SNP_2": "S1"}

    def test_invalid_risk_level_degrades_to_none_without_raising(self, workbook_path: Path) -> None:
        rows = [
            _row(item_id="SNP_1", risk_level=99, requires_human_review=False),
            _row(item_id="SNP_2", risk_level=4, requires_human_review=False),
        ]
        _write_workbook(workbook_path, rows)
        claims = CorpusClaimDataset(path=workbook_path).load()
        by_id = {c.claim_id: c.risk_level for c in claims}
        assert by_id == {"SNP_1": None, "SNP_2": 4}

    def test_boolean_risk_level_does_not_get_miscast_as_one(self, workbook_path: Path) -> None:
        """
        Regression test: bool is a subclass of int in Python, so a stray
        boolean cell (True == 1) could otherwise slip through an
        `isinstance(x, int)` + range check and be silently miscast as
        risk_level=1.
        """
        _write_workbook(
            workbook_path,
            [_row(risk_level=True, requires_human_review=False)],
        )
        claim = CorpusClaimDataset(path=workbook_path).load()[0]
        assert claim.risk_level is None

    def test_workbook_missing_claim_original_column_yields_no_claims(
        self, workbook_path: Path
    ) -> None:
        """
        A column entirely absent from the header means every row's
        ``claim_original`` reads back as None via ``.get()`` — the same
        blank-row skip that handles an individual blank cell (see
        TestCorpusClaimDatasetEthicalGate) also covers this case, so the
        whole file degrades gracefully to zero claims rather than raising.
        """
        wb_rows = [_row(requires_human_review=False)]
        header_without_claim_original = tuple(c for c in _HEADER if c != "claim_original")
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "01_Corpus_claim"
        sheet.append(header_without_claim_original)
        for row in wb_rows:
            sheet.append([row.get(col) for col in header_without_claim_original])
        wb.save(workbook_path)
        assert CorpusClaimDataset(path=workbook_path).load() == []

    def test_respects_max_claims(self, workbook_path: Path) -> None:
        rows = [
            _row(item_id=f"SNP_{i}", requires_human_review=False) for i in range(5)
        ]
        _write_workbook(workbook_path, rows)
        claims = CorpusClaimDataset(path=workbook_path).load(max_claims=2)
        assert len(claims) == 2


# ─── Errors ───────────────────────────────────────────────────────────────────

class TestCorpusClaimDatasetErrors:
    def test_load_raises_on_missing_file(self, workbook_path: Path) -> None:
        with pytest.raises(IngestionError, match="Could not read"):
            CorpusClaimDataset(path=workbook_path).load()

    def test_load_raises_on_missing_sheet(self, workbook_path: Path) -> None:
        wb = openpyxl.Workbook()
        wb.active.title = "SomeOtherSheet"
        wb.save(workbook_path)
        with pytest.raises(IngestionError, match="01_Corpus_claim"):
            CorpusClaimDataset(path=workbook_path).load()

    def test_load_raises_on_corrupt_file(self, workbook_path: Path) -> None:
        """A .xlsx-named file that isn't actually a zip archive at all."""
        workbook_path.write_text("this is not a real xlsx file", encoding="utf-8")
        with pytest.raises(IngestionError, match="Could not read"):
            CorpusClaimDataset(path=workbook_path).load()

    def test_load_raises_on_path_that_is_a_directory(self, workbook_path: Path) -> None:
        workbook_path.mkdir()
        with pytest.raises(IngestionError, match="Could not read"):
            CorpusClaimDataset(path=workbook_path).load()

    def test_load_raises_on_duplicate_header_columns(self, workbook_path: Path) -> None:
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "01_Corpus_claim"
        duplicated_header = _HEADER + ("item_id",)
        sheet.append(duplicated_header)
        wb.save(workbook_path)
        with pytest.raises(IngestionError, match="duplicate column"):
            CorpusClaimDataset(path=workbook_path).load()

    def test_load_handles_row_shorter_than_header_without_raising(
        self, workbook_path: Path
    ) -> None:
        """
        Some non-Excel export tools omit the sheet's <dimension> element,
        which can make openpyxl's read_only mode yield trailing-cell-
        truncated rows. Missing trailing cells must degrade to None like
        an explicitly blank cell, not raise IndexError.
        """
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "01_Corpus_claim"
        sheet.append(_HEADER)
        full_row = _row(requires_human_review=False)
        # Only append values for the first half of the columns.
        truncated_values = [full_row.get(col) for col in _HEADER[: len(_HEADER) // 2]]
        sheet.append(truncated_values)
        wb.save(workbook_path)
        # include_unreviewed=True: requires_human_review lives in the
        # truncated-away second half of the header, so it pads to None —
        # the point of this test is only that padding doesn't raise
        # IndexError, not the (separate, already-tested) ethical gate.
        claims = CorpusClaimDataset(path=workbook_path).load(include_unreviewed=True)
        assert len(claims) == 1


# ─── Internal edge cases (mocked openpyxl) ────────────────────────────────────

def _mock_workbook(rows_after_header: list[Any], header: tuple = _HEADER) -> MagicMock:
    mock_sheet = MagicMock()
    mock_sheet.iter_rows.return_value = iter([header, *rows_after_header])
    mock_workbook = MagicMock()
    mock_workbook.sheetnames = ["01_Corpus_claim"]
    mock_workbook.__getitem__.return_value = mock_sheet
    return mock_workbook


class TestCorpusClaimDatasetInternalEdgeCases:
    def test_load_raises_actionable_error_when_openpyxl_not_installed(
        self, workbook_path: Path
    ) -> None:
        with patch.dict(sys.modules, {"openpyxl": None}):
            with pytest.raises(IngestionError, match="openpyxl is required"):
                CorpusClaimDataset(path=workbook_path).load()

    def test_load_returns_empty_list_when_sheet_has_no_header_row(
        self, workbook_path: Path
    ) -> None:
        """Covers _read_rows' `except StopIteration: return []` branch."""
        mock_workbook = MagicMock()
        mock_workbook.sheetnames = ["01_Corpus_claim"]
        mock_sheet = MagicMock()
        mock_sheet.iter_rows.return_value = iter([])  # no header row at all
        mock_workbook.__getitem__.return_value = mock_sheet
        with patch("openpyxl.load_workbook", return_value=mock_workbook):
            claims = CorpusClaimDataset(path=workbook_path).load()
        assert claims == []
        mock_workbook.close.assert_called_once()

    def test_load_skips_a_none_row_yielded_by_iter_rows(self, workbook_path: Path) -> None:
        """
        Covers the `if raw_row is None: continue` branch in _read_rows —
        openpyxl's read_only mode can yield a bare None for certain
        sparse/blank rows rather than a tuple of None cell values.
        """
        valid_row = tuple(_row(requires_human_review=False).get(col) for col in _HEADER)
        mock_workbook = _mock_workbook([None, valid_row])
        with patch("openpyxl.load_workbook", return_value=mock_workbook):
            claims = CorpusClaimDataset(path=workbook_path).load()
        assert len(claims) == 1

    def test_read_rows_closes_workbook_even_on_error(self, workbook_path: Path) -> None:
        mock_workbook = MagicMock()
        mock_workbook.sheetnames = ["SomeOtherSheet"]  # triggers the missing-sheet error
        with patch("openpyxl.load_workbook", return_value=mock_workbook):
            with pytest.raises(IngestionError, match="01_Corpus_claim"):
                CorpusClaimDataset(path=workbook_path).load()
        mock_workbook.close.assert_called_once()


# ─── content_hash ─────────────────────────────────────────────────────────────

class TestCorpusClaimDatasetContentHash:
    def test_content_hash_before_load_is_all_zeros(self, workbook_path: Path) -> None:
        assert CorpusClaimDataset(path=workbook_path).content_hash == "0" * 16

    def test_content_hash_after_load_is_deterministic(self, workbook_path: Path) -> None:
        _write_workbook(workbook_path, [_row(requires_human_review=False)])
        first = CorpusClaimDataset(path=workbook_path)
        first.load()
        second = CorpusClaimDataset(path=workbook_path)
        second.load()
        assert first.content_hash == second.content_hash
        assert len(first.content_hash) == 16


# ─── download() ─────────────────────────────────────────────────────────────

class TestCorpusClaimDatasetDownload:
    def test_download_noops_when_file_already_present(self, tmp_path: Path) -> None:
        _write_workbook(tmp_path / "Corpus_claim_RAG_Mistral_output_v3.xlsx", [_row()])
        CorpusClaimDataset().download(str(tmp_path))  # must not raise

    def test_download_raises_when_file_missing(self, tmp_path: Path) -> None:
        with pytest.raises(IngestionError, match="no automated download"):
            CorpusClaimDataset().download(str(tmp_path))
