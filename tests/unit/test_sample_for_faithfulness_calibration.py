"""
Unit tests for scripts/sample_for_faithfulness_calibration.py.

Not covered by the coverage gate (pyproject.toml omits scripts/*), but the
stratification and ground-truth-lookup logic is non-trivial enough to
warrant regression tests, matching test_sample_corpus_claim_for_review.py's
own scope note.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest

from eiger.core.models import (
    Document,
    EvaluationRecord,
    GenerationResult,
    RetrievalResult,
    RetrievedDocument,
)

_SCRIPT_PATH = (
    Path(__file__).resolve().parents[2] / "scripts" / "sample_for_faithfulness_calibration.py"
)


@pytest.fixture(scope="module")
def calib_module():
    spec = importlib.util.spec_from_file_location("sample_for_faithfulness_calibration", _SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _record(claim_id: str, poisoned: bool, gt_text: str | None = "the ground truth fact") -> EvaluationRecord:
    hits = []
    if gt_text is not None:
        hits.append(
            RetrievedDocument(
                document=Document(doc_id=f"gt_{claim_id}", claim_id=claim_id, text=gt_text, doc_type="ground_truth"),
                score=0.9,
                rank=1,
            )
        )
    if poisoned:
        hits.append(
            RetrievedDocument(
                document=Document(doc_id=f"poison_{claim_id}", claim_id=claim_id, text="poisoned text", doc_type="poisoned"),
                score=0.8,
                rank=len(hits) + 1,
            )
        )
    retrieval = RetrievalResult(query=f"query for {claim_id}", claim_id=claim_id, hits=hits, top_k=5)
    generation = GenerationResult(
        claim_id=claim_id, query=retrieval.query, context_docs=[h.document.text for h in hits],
        answer=f"answer for {claim_id}", model_name="test-model",
    )
    return EvaluationRecord(
        claim_id=claim_id, generation=generation, retrieval=retrieval,
        metrics={"ragas_faithfulness": 0.5, "ragas_answer_correctness": 0.6},
    )


class TestGroundTruthText:
    def test_returns_ground_truth_doc_text_when_retrieved(self, calib_module) -> None:
        record = _record("C1", poisoned=False, gt_text="the real fact")
        assert calib_module._ground_truth_text(record) == "the real fact"

    def test_returns_none_when_ground_truth_not_retrieved(self, calib_module) -> None:
        record = _record("C1", poisoned=True, gt_text=None)
        assert calib_module._ground_truth_text(record) is None

    def test_ignores_ground_truth_doc_belonging_to_a_different_claim(self, calib_module) -> None:
        """A hit with doc_type=ground_truth but a mismatched claim_id must not be used."""
        other_gt = RetrievedDocument(
            document=Document(doc_id="gt_other", claim_id="OTHER", text="not this claim's fact", doc_type="ground_truth"),
            score=0.9, rank=1,
        )
        retrieval = RetrievalResult(query="q", claim_id="C1", hits=[other_gt], top_k=5)
        generation = GenerationResult(claim_id="C1", query="q", context_docs=["not this claim's fact"], answer="a", model_name="m")
        record = EvaluationRecord(claim_id="C1", generation=generation, retrieval=retrieval)
        assert calib_module._ground_truth_text(record) is None


class TestSelectSample:
    def test_covers_both_strata_when_both_present(self, calib_module) -> None:
        records = [_record(f"P{i}", poisoned=True) for i in range(20)] + [_record(f"C{i}", poisoned=False) for i in range(20)]
        sampled = calib_module.select_sample(records, sample_size=10, seed=42)
        assert any(r.retrieval.contains_poisoned for r in sampled)
        assert any(not r.retrieval.contains_poisoned for r in sampled)

    def test_deterministic_for_same_seed(self, calib_module) -> None:
        records = [_record(f"C{i}", poisoned=(i % 2 == 0)) for i in range(30)]
        first = calib_module.select_sample(records, sample_size=10, seed=7)
        second = calib_module.select_sample(records, sample_size=10, seed=7)
        assert [r.claim_id for r in first] == [r.claim_id for r in second]

    def test_single_stratum_still_works(self, calib_module) -> None:
        records = [_record(f"C{i}", poisoned=False) for i in range(5)]
        sampled = calib_module.select_sample(records, sample_size=3, seed=42)
        assert len(sampled) == 3
        assert all(not r.retrieval.contains_poisoned for r in sampled)

    def test_empty_records_returns_empty(self, calib_module) -> None:
        assert calib_module.select_sample([], sample_size=10, seed=42) == []


class TestWriteAnnotationWorkbook:
    def test_writes_workbook_and_returns_automated_scores(self, calib_module, tmp_path) -> None:
        records = [_record("C1", poisoned=False), _record("C2", poisoned=True)]
        output_path = tmp_path / "batch.xlsx"

        scores = calib_module.write_annotation_workbook(records, output_path)

        assert output_path.exists()
        assert scores == {
            "C1": {"ragas_faithfulness": 0.5, "ragas_answer_correctness": 0.6},
            "C2": {"ragas_faithfulness": 0.5, "ragas_answer_correctness": 0.6},
        }

    def test_workbook_never_contains_automated_scores_in_visible_columns(self, calib_module, tmp_path) -> None:
        """Blind-annotation guarantee: the numeric automated scores must not leak into the sheet."""
        import openpyxl

        records = [_record("C1", poisoned=False)]
        output_path = tmp_path / "batch.xlsx"
        calib_module.write_annotation_workbook(records, output_path)

        wb = openpyxl.load_workbook(output_path)
        header = [c.value for c in wb.active[1]]
        assert "ragas_faithfulness" not in header
        assert "ragas_answer_correctness" not in header

    def test_missing_ground_truth_gets_placeholder_text(self, calib_module, tmp_path) -> None:
        records = [_record("C1", poisoned=True, gt_text=None)]
        output_path = tmp_path / "batch.xlsx"
        calib_module.write_annotation_workbook(records, output_path)

        import openpyxl

        wb = openpyxl.load_workbook(output_path)
        header = [c.value for c in wb.active[1]]
        row = [c.value for c in wb.active[2]]
        gt_col = header.index("ground_truth_fact")
        assert "not retrieved" in row[gt_col]


class TestMainCLI:
    def test_missing_input_file_returns_1(self, calib_module, tmp_path, capsys) -> None:
        exit_code = calib_module.main([str(tmp_path / "nope.json"), "-o", str(tmp_path / "out.xlsx")])
        assert exit_code == 1
        assert "not found" in capsys.readouterr().err

    def test_empty_records_returns_0_without_writing(self, calib_module, tmp_path) -> None:
        results_path = tmp_path / "results.json"
        results_path.write_text(
            '{"experiment_id": "e", "config_hash": "h", "config": '
            '{"dataset": {"name": "json_fixture"}, "attacks": [], '
            '"retriever": {"type": "sparse"}, "llm": {"backend": "ollama", "model": "m"}, '
            '"metrics": [], "output_dir": "results/e/"}, "records": []}'
        )
        output_path = tmp_path / "out.xlsx"
        exit_code = calib_module.main([str(results_path), "-o", str(output_path)])
        assert exit_code == 0
        assert not output_path.exists()

    def test_end_to_end_writes_workbook_and_reference_file(self, calib_module, tmp_path) -> None:
        from eiger.core.models import ExperimentConfig, ExperimentResult

        records = [_record(f"C{i}", poisoned=(i % 2 == 0)) for i in range(6)]
        config = ExperimentConfig(
            dataset={"name": "json_fixture"}, attacks=[],
            retriever={"type": "sparse"}, llm={"backend": "ollama", "model": "m"},
            metrics=[], output_dir=str(tmp_path / "results"),
        )
        result = ExperimentResult(experiment_id="e", config_hash="h", config=config, records=records)
        results_path = tmp_path / "results.json"
        results_path.write_text(result.to_json(), encoding="utf-8")

        output_path = tmp_path / "batch.xlsx"
        exit_code = calib_module.main([str(results_path), "-o", str(output_path), "--sample-size", "4"])

        assert exit_code == 0
        assert output_path.exists()
        reference_path = output_path.with_suffix(output_path.suffix + ".automated_scores.json")
        assert reference_path.exists()
